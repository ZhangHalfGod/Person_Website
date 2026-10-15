"""Mark Goal3 rows in 房间烘焙与后续开发计划.xlsx. Close Excel first."""
import openpyxl

path = r"d:\trae_git\Person_Website\System_doc\3D美术\第二期\房间烘焙与后续开发计划.xlsx"
wb = openpyxl.load_workbook(path)
ws = wb[wb.sheetnames[3]]
print("sheet", repr(ws.title))

ws.cell(2, 2).value = "进行中（主项已调）"

# Col layout: A=ID, E=状态 (1-indexed: 5)
done = {
    "G3-01": "已完成",
    "G3-02": "已完成",
    "G3-03": "已完成",  # floor local rebake + polish; residual grain noted in Bake2 md
    "G3-04": "已完成",
    "G3-05": "未开始",  # gold-frame carving not done; no fake check
    "G3-06": "已完成",
}
skip_ibl = {
    "G3-IBL-01": "跳过",
    "G3-IBL-02": "跳过",
}

for row in range(9, 22):
    tid = ws.cell(row, 1).value
    if tid in done:
        ws.cell(row, 5).value = done[tid]
        print("set", tid, done[tid])
    elif tid in skip_ibl:
        ws.cell(row, 5).value = skip_ibl[tid]
        print("set", tid, skip_ibl[tid])

wb.save(path)
print("saved")

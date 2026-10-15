import openpyxl

path = r"d:\trae_git\Person_Website\System_doc\3D美术\第二期\房间烘焙与后续开发计划.xlsx"
wb = openpyxl.load_workbook(path)
ws = wb[wb.sheetnames[2]]
print("sheet", repr(ws.title))
ws.cell(2, 2).value = "进行中（主项已落地）"
done = {
    "B1-01",
    "B1-02",
    "B1-03",
    "B1-04",
    "B1-05",
    "B1-06",
    "B1-07",
    "B1-08",
    "B1-09",
    "B1-10",
}
for row in range(9, 22):
    tid = ws.cell(row, 1).value
    if tid in done:
        ws.cell(row, 6).value = "已完成"
        print("done", tid)
    elif tid == "B1-11":
        print("leave", tid, ws.cell(row, 6).value)
wb.save(path)
print("saved")
